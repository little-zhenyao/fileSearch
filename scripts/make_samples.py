"""Generate reproducible test materials; these are synthetic, not real course files."""
from __future__ import annotations

import argparse
from pathlib import Path

import pymupdf
from docx import Document


def generate(destination: Path):
    destination.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((60, 70), 'CS2106 | Operating Systems', fontsize=11, color=(.4,.5,.4))
    page.insert_text((60, 120), '07  Deadlocks', fontsize=28, color=(.15,.3,.22))
    lines = [
        'A deadlock occurs when a set of processes cannot make progress.',
        'Each process waits for a resource held by another process.',
        '', 'Four necessary conditions', '',
        '1. Mutual exclusion: a resource can be held by only one process.',
        '2. Hold and wait: processes hold resources while requesting more.',
        '3. No preemption: resources cannot be forcibly removed.',
        '4. Circular wait: a closed chain of waiting processes exists.',
        '', 'Deadlock prevention breaks at least one of these conditions.',
        'Deadlocks is a different word form and is not expanded automatically.',
    ]
    page.insert_text((60, 165), '\n'.join(lines), fontsize=11, lineheight=1.8)
    page.insert_text((60, 730), 'Synthetic acceptance fixture | page 1', fontsize=9)
    page = pdf.new_page(width=612, height=792)
    page.insert_text((60, 90), 'Resource allocation graphs', fontsize=23)
    page.insert_text((60, 150), 'A cycle may indicate deadlock.\nThe graph helps explain mutual exclusion.\nRepeated term: deadlock. deadlock.', fontsize=12, lineheight=1.8)
    page = pdf.new_page(width=612, height=792)
    page.draw_rect(pymupdf.Rect(70, 100, 540, 650), color=(.8,.85,.8), fill=(.97,.98,.96))
    pdf.save(destination / 'Lecture07_Deadlocks.pdf'); pdf.close()

    for angle in (0, 90, 180, 270):
        pdf = pymupdf.open(); page = pdf.new_page(width=600, height=800)
        page.insert_text((100, 150), f'Deadlock: cropped page at {angle} degrees.', fontsize=16)
        page.insert_text((100, 200), 'dead-\nlock and mutual exclusion', fontsize=14)
        page.set_cropbox(pymupdf.Rect(40, 60, 560, 740)); page.set_rotation(angle)
        pdf.save(destination / f'Crop_Rotation_{angle}.pdf'); pdf.close()

    pdf = pymupdf.open(); page = pdf.new_page(width=612, height=792)
    page.insert_text((50, 90), 'Left column: deadlock.\nMutual exclusion is required.', fontsize=11)
    page.insert_text((330, 90), 'Right column: starvation.\nIndefinite blocking.', fontsize=11)
    pdf.save(destination / 'Two_Columns.pdf'); pdf.close()
    pdf = pymupdf.open(); page = pdf.new_page()
    page.insert_text((60, 100), '操作系统中的死锁与资源等待。', fontname='china-s', fontsize=18)
    page.insert_text((60, 150), 'deadlock / 中文混合检索', fontname='china-s', fontsize=16)
    pdf.save(destination / 'Mixed_Chinese_English.pdf'); pdf.close()
    pdf = pymupdf.open(); pdf.new_page(); pdf.save(destination / 'Scan_No_Text.pdf'); pdf.close()
    pdf = pymupdf.open(); page = pdf.new_page(); page.insert_text((60, 100), 'Protected example')
    pdf.save(destination / 'Password_Required.pdf', encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='example'); pdf.close()
    (destination / 'Corrupt_File.pdf').write_bytes(b'This is intentionally not a PDF.')

    document = Document()
    document.add_heading('CS2106 操作系统速记', level=1)
    document.add_paragraph('Deadlock requires mutual exclusion. 中文死锁与资源等待。')
    document.add_heading('Concept comparison', level=2)
    table = document.add_table(rows=3, cols=2)
    table.cell(0,0).text = 'Concept'; table.cell(0,1).text = 'Definition'
    table.cell(1,0).text = 'Starvation'; table.cell(1,1).text = 'Indefinite blocking: a process waits without progress.'
    table.cell(2,0).text = 'Deadlock'; table.cell(2,1).text = 'Circular wait and mutual exclusion. 资源等待示例。'
    document.add_paragraph('Unicode offsets: 😀 The ﬃ ligature, café and deadlock.')
    document.save(destination / 'Summary_Concepts.docx')
    document = Document(); document.add_heading('Tutorial 07', level=1)
    for index in range(30):
        document.add_paragraph(f'Exercise {index + 1}: Explain deadlock prevention and mutual exclusion.')
    document.save(destination / 'Tutorial07.docx')
    print(f'Generated 12 synthetic PDF/DOCX samples at {destination.resolve()}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--output', type=Path, default=Path('.local/samples'))
    generate(parser.parse_args().output)
