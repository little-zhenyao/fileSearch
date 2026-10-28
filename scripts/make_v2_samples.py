"""Generate synthetic navigation/dual-reader fixtures outside the user's library."""
from pathlib import Path
import pymupdf
from docx import Document


def generate(destination):
    destination.mkdir(parents=True, exist_ok=True)
    for name, pages in [('V2_Lecture.pdf', 30), ('V2_Tutorial.pdf', 3)]:
        document = pymupdf.open()
        for index in range(pages):
            page = document.new_page(width=600, height=800)
            lines = [f'Operating systems - page {index + 1}', 'Deadlock requires mutual exclusion.',
                     'Deadlocks and resource allocation.', 'Starvation means indefinite blocking.']
            lines.extend([f'Example {line}: processes and threads share resources.' for line in range(25)])
            page.insert_text((55, 75), '\n'.join(lines), fontsize=12, lineheight=1.5)
        document.save(destination / name)
        document.close()
    for name in ['V2_Notes.docx', 'V2_Exercises.docx']:
        document = Document()
        document.add_heading(name + ' Operating Systems 操作系统', level=1)
        for index in range(35):
            document.add_paragraph(f'Topic {index + 1}: Deadlock and resource allocation. 中文死锁与资源等待。')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Concept'; table.cell(0, 1).text = 'Definition'
        table.cell(1, 0).text = 'Starvation'; table.cell(1, 1).text = 'Indefinite blocking and 资源等待'
        document.save(destination / name)
    print(destination.resolve())


if __name__ == '__main__':
    generate(Path('.local/v2-qa/fixtures'))
