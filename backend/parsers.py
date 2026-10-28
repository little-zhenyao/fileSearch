"""Extract original text, stable DOCX anchors, and PDF user-space rectangles."""
from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


class PasswordRequired(Exception):
    pass


def parse_pdf(path: Path) -> tuple[list[dict], str | None]:
    units = []
    with pymupdf.open(path) as document:
        if document.needs_pass:
            raise PasswordRequired('文件需要密码；请先在外部生成允许使用的未加密副本。')
        for number, page in enumerate(document, 1):
            # get_text coordinates are unrotated, crop-relative MuPDF coordinates.
            # PDF.js expects native PDF user space; remove rotation before obtaining
            # transformation_matrix (MuPDF otherwise changes its crop translation).
            rotation = page.rotation
            page.set_rotation(0)
            inverse = ~page.transformation_matrix
            chars, boxes, line_ids = [], [], []
            raw = page.get_text('rawdict', sort=True, flags=pymupdf.TEXTFLAGS_RAWDICT & ~pymupdf.TEXT_PRESERVE_IMAGES)
            line_id = 0
            for block in raw['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        for char in span['chars']:
                            rect = pymupdf.Rect(char['bbox']) * inverse
                            for codepoint in char['c']:
                                chars.append(codepoint)
                                boxes.append([round(v, 3) for v in rect])
                                line_ids.append(line_id)
                    chars.append('\n')
                    boxes.append(None)
                    line_ids.append(line_id)
                    line_id += 1
            page.set_rotation(rotation)
            text = ''.join(chars)
            units.append({
                'text': text, 'label': f'第 {number} 页',
                'anchor': {'page': number, 'boxes': boxes, 'lines': line_ids,
                           'rotation': rotation, 'type': 'page'},
            })
    return units, '图片、公式和无文字层页面不参与文字检索。'


def parse_docx(path: Path) -> tuple[list[dict], str | None]:
    document = Document(path)
    units: list[dict] = []
    paragraph_number = 0
    table_number = 0
    section = ''

    def table_units(table: Table, key: str, label: str):
        # Skip repeated references to merged cells, retaining the physical grid span.
        seen = set()
        for row_index, row in enumerate(table.rows):
            column = 0
            for cell in row.cells:
                identity = cell._tc
                if identity in seen:
                    column += 1
                    continue
                seen.add(identity)
                cell_key = f'{key}-r{row_index + 1}-c{column + 1}'
                text = '\n'.join(p.text for p in cell.paragraphs)
                units.append({'text': text,
                              'label': f'{label} · 行 {row_index + 1} 列 {column + 1}',
                              'anchor': {'type': 'cell', 'key': cell_key, 'table': key,
                                         'row': row_index, 'column': column,
                                         'colspan': cell._tc.grid_span, 'section': section}})
                for nested_index, nested in enumerate(cell.tables, 1):
                    table_units(nested, f'{cell_key}-t{nested_index}', f'{label} 嵌套表 {nested_index}')
                column += 1

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            style = block.style.name if block.style else ''
            heading = style.startswith('Heading') or style == 'Title'
            if heading:
                section = block.text
            units.append({'text': block.text,
                          'label': f'{section + " · " if section else ""}段落 {paragraph_number}',
                          'anchor': {'type': 'paragraph', 'key': f'p{paragraph_number}',
                                     'heading': heading, 'section': section}})
        elif isinstance(block, Table):
            table_number += 1
            table_units(block, f't{table_number}', f'表格 {table_number}')
    return units, 'DOCX 为结构化文本预览，无 Word 页码；图片、浮动对象、文本框、页眉页脚和批注不参与检索。'


def parse(path: Path, kind: str):
    return parse_pdf(path) if kind == 'pdf' else parse_docx(path)


def rectangles(anchor: dict, ranges: list[list[int]]) -> list[list[list[float]]]:
    """Merge adjacent characters on each line; keep separated occurrences separate."""
    result = []
    for start, end in ranges:
        by_line: dict[int, list[float]] = {}
        for i in range(start, min(end, len(anchor.get('boxes', [])))):
            rect = anchor['boxes'][i]
            if rect is None:
                continue
            line = anchor['lines'][i]
            if line in by_line:
                old = by_line[line]
                old[:] = [min(old[0], rect[0]), min(old[1], rect[1]),
                          max(old[2], rect[2]), max(old[3], rect[3])]
            else:
                by_line[line] = list(rect)
        result.append(list(by_line.values()))
    return result
