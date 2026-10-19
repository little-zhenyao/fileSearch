from __future__ import annotations

import io
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend.app import create_app


@pytest.fixture
def app(tmp_path):
    return create_app(config_path=tmp_path / 'settings.json', library_path=tmp_path / 'library', start_worker=False)


@pytest.fixture
def client(app):
    with TestClient(app, headers={'X-Local-Request': '1'}) as client:
        yield client


@pytest.fixture
def course(client):
    response = client.post('/api/courses', json={'name': 'CS2106 操作系统'})
    assert response.status_code == 201
    return response.json()['id']


def pdf_bytes(text='Deadlock requires mutual exclusion. Deadlock differs from deadlocks.', blank=False):
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_text((72, 120), text, fontsize=14)
    if blank:
        document.new_page(width=600, height=800)
    data = document.tobytes()
    document.close()
    return data


def docx_bytes(text='Deadlock requires mutual exclusion. 中文死锁与资源等待。'):
    document = Document()
    document.add_heading('Operating Systems 操作系统', level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Concept'
    table.cell(0, 1).text = 'Definition'
    table.cell(1, 0).text = 'Starvation'
    table.cell(1, 1).text = 'Indefinite blocking and 资源等待 are related.'
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def import_bytes(client, app, course, files):
    import json
    response = client.post('/api/import-jobs', data={'course_id': course, 'paths': json.dumps([name for name, _ in files])},
                           files=[('files', (name, data, 'application/octet-stream')) for name, data in files])
    assert response.status_code == 202, response.text
    while app.state.library.process_next():
        pass
    return client.get('/api/import-jobs/' + response.json()['job_id']).json()
